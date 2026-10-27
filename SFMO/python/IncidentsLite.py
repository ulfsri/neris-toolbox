import sys
import subprocess
import os
import csv
import re
import json
import time
import pathlib
import threading
import traceback
from collections import deque
from datetime import datetime, timedelta, timezone
from urllib.parse import urlsplit, parse_qs
from concurrent.futures import ThreadPoolExecutor, as_completed


def ensure_dependencies():
    def pip_install(*packages):
        subprocess.run(
            [sys.executable, "-m", "pip", "install", *packages, "--quiet"],
            check=True
        )
    try:
        import pandas
        import openpyxl
    except ImportError:
        print("Installing pandas and openpyxl...")
        pip_install("pandas", "openpyxl")
        print("✓ pandas and openpyxl installed")

    try:
        from neris_api_client import NerisApiClient
        print("✓ Dependencies ready")
    except ImportError:
        print("Installing neris-api-client...")
        pip_install("neris-api-client")
        print("✓ neris-api-client installed")


DOCS_DIR = pathlib.Path.home() / "Documents"


INCIDENT_TYPE_MAP = {
    "FIRE – Outside Fire": [
        "FIRE||OUTSIDE_FIRE||CONSTRUCTION_WASTE",
        "FIRE||OUTSIDE_FIRE||DUMPSTER_OUTDOOR_CONTAINER_FIRE",
        "FIRE||OUTSIDE_FIRE||OTHER_OUTSIDE_FIRE",
        "FIRE||OUTSIDE_FIRE||OUTSIDE_TANK_FIRE",
        "FIRE||OUTSIDE_FIRE||TRASH_RUBBISH_FIRE",
        "FIRE||OUTSIDE_FIRE||UTILITY_INFRASTRUCTURE_FIRE",
        "FIRE||OUTSIDE_FIRE||VEGETATION_GRASS_FIRE",
        "FIRE||OUTSIDE_FIRE||WILDFIRE_URBAN_INTERFACE",
        "FIRE||OUTSIDE_FIRE||WILDFIRE_WILDLAND",
    ],
    "FIRE – Special Fire": [
        "FIRE||SPECIAL_FIRE||ESS_FIRE",
        "FIRE||SPECIAL_FIRE||EXPLOSION",
        "FIRE||SPECIAL_FIRE||INFRASTRUCTURE_FIRE",
    ],
    "FIRE – Structure Fire": [
        "FIRE||STRUCTURE_FIRE||CHIMNEY_FIRE",
        "FIRE||STRUCTURE_FIRE||CONFINED_COOKING_APPLIANCE_FIRE",
        "FIRE||STRUCTURE_FIRE||ROOM_AND_CONTENTS_FIRE",
        "FIRE||STRUCTURE_FIRE||STRUCTURAL_INVOLVEMENT_FIRE",
    ],
    "FIRE – Transportation Fire": [
        "FIRE||TRANSPORTATION_FIRE||AIRCRAFT_FIRE",
        "FIRE||TRANSPORTATION_FIRE||BOAT_PERSONAL_WATERCRAFT_BARGE_FIRE",
        "FIRE||TRANSPORTATION_FIRE||POWERED_MOBILITY_DEVICE_FIRE",
        "FIRE||TRANSPORTATION_FIRE||TRAIN_RAIL_FIRE",
        "FIRE||TRANSPORTATION_FIRE||VEHICLE_FIRE_COMMERCIAL",
        "FIRE||TRANSPORTATION_FIRE||VEHICLE_FIRE_FOOD_TRUCK",
        "FIRE||TRANSPORTATION_FIRE||VEHICLE_FIRE_PASSENGER",
        "FIRE||TRANSPORTATION_FIRE||VEHICLE_FIRE_RV",
    ],
    "HAZSIT – Hazardous Materials": [
        "HAZSIT||HAZARDOUS_MATERIALS||BIOLOGICAL_RELEASE_INCIDENT",
        "HAZSIT||HAZARDOUS_MATERIALS||CARBON_MONOXIDE_RELEASE",
        "HAZSIT||HAZARDOUS_MATERIALS||FUEL_SPILL_ODOR",
        "HAZSIT||HAZARDOUS_MATERIALS||GAS_LEAK_ODOR",
        "HAZSIT||HAZARDOUS_MATERIALS||HAZMAT_RELEASE_FACILITY",
        "HAZSIT||HAZARDOUS_MATERIALS||HAZMAT_RELEASE_TRANSPORT",
        "HAZSIT||HAZARDOUS_MATERIALS||RADIOACTIVE_RELEASE_INCIDENT",
    ],
    "HAZSIT – Non-Chemical Hazard": [
        "HAZSIT||HAZARD_NONCHEM||BOMB_THREAT_RESPONSE_SUSPICIOUS_PACKAGE",
        "HAZSIT||HAZARD_NONCHEM||ELEC_HAZARD_SHORT_CIRCUIT",
        "HAZSIT||HAZARD_NONCHEM||ELEC_POWER_LINE_DOWN_ARCHING_MALFUNC",
        "HAZSIT||HAZARD_NONCHEM||MOTOR_VEHICLE_COLLISION",
    ],
    "HAZSIT – Investigation": [
        "HAZSIT||INVESTIGATION||ODOR",
        "HAZSIT||INVESTIGATION||SMOKE_INVESTIGATION",
    ],
    "HAZSIT – Overpressure": [
        "HAZSIT||OVERPRESSURE||NO_RUPTURE",
        "HAZSIT||OVERPRESSURE||RUPTURE_WITHOUT_FIRE",
    ],
    "LAWENFORCE": ["LAWENFORCE"],
    "MEDICAL – Illness": [
        "MEDICAL||ILLNESS", "MEDICAL||ILLNESS||ABDOMINAL_PAIN",
        "MEDICAL||ILLNESS||ALLERGIC_REACTION_STINGS", "MEDICAL||ILLNESS||ALTERED_MENTAL_STATUS",
        "MEDICAL||ILLNESS||BACK_PAIN_NON_TRAUMA", "MEDICAL||ILLNESS||BREATHING_PROBLEMS",
        "MEDICAL||ILLNESS||CARDIAC_ARREST", "MEDICAL||ILLNESS||CHEST_PAIN_NON_TRAUMA",
        "MEDICAL||ILLNESS||CONVULSIONS_SEIZURES", "MEDICAL||ILLNESS||DIABETIC_PROBLEMS",
        "MEDICAL||ILLNESS||HEADACHE", "MEDICAL||ILLNESS||HEART_PROBLEMS",
        "MEDICAL||ILLNESS||NAUSEA_VOMITING", "MEDICAL||ILLNESS||NO_APPROPRIATE_CHOICE",
        "MEDICAL||ILLNESS||OVERDOSE", "MEDICAL||ILLNESS||PANDEMIC_EPIDEMIC_OUTBREAK",
        "MEDICAL||ILLNESS||PREGNANCY_CHILDBIRTH", "MEDICAL||ILLNESS||PSYCHOLOGICAL_BEHAVIOR_ISSUES",
        "MEDICAL||ILLNESS||SICK_CASE", "MEDICAL||ILLNESS||STROKE_CVA",
        "MEDICAL||ILLNESS||UNCONSCIOUS_VICTIM", "MEDICAL||ILLNESS||UNKNOWN_PROBLEM",
        "MEDICAL||ILLNESS||WELL_PERSON_CHECK",
    ],
    "MEDICAL – Injury": [
        "MEDICAL||INJURY", "MEDICAL||INJURY||ANIMAL_BITES", "MEDICAL||INJURY||ASSAULT",
        "MEDICAL||INJURY||BURNS_EXPLOSION", "MEDICAL||INJURY||CARBON_MONOXIDE_OTHER_INHALATION_INJURY",
        "MEDICAL||INJURY||CHOKING", "MEDICAL||INJURY||DROWNING_DIVING_SCUBA_ACCIDENT",
        "MEDICAL||INJURY||ELECTROCUTION", "MEDICAL||INJURY||EYE_TRAUMA",
        "MEDICAL||INJURY||FALL", "MEDICAL||INJURY||GUNSHOT_WOUND",
        "MEDICAL||INJURY||HEAT_COLD_EXPOSURE", "MEDICAL||INJURY||HEMORRHAGE_LACERATION",
        "MEDICAL||INJURY||INDUSTRIAL_INACCESSIBLE_ENTRAPMENT",
        "MEDICAL||INJURY||MOTOR_VEHICLE_COLLISION", "MEDICAL||INJURY||OTHER_TRAUMATIC_INJURY",
        "MEDICAL||INJURY||POISONING", "MEDICAL||INJURY||STAB_PENETRATING_TRAUMA",
    ],
    "MEDICAL – Other": [
        "MEDICAL||OTHER||AIRMEDICAL_TRANSPORT", "MEDICAL||OTHER||COMMUNITY_PUBLIC_HEALTH",
        "MEDICAL||OTHER||HEALTHCARE_PROFESSIONAL_ADMISSION", "MEDICAL||OTHER||INTERCEPT_OTHER_UNIT",
        "MEDICAL||OTHER||MEDICAL_ALARM", "MEDICAL||OTHER||STANDBY_REQUEST",
        "MEDICAL||OTHER||TRANSFER_INTERFACILITY",
    ],
    "NOEMERG – Cancelled": ["NOEMERG||CANCELLED"],
    "NOEMERG – False Alarm": [
        "NOEMERG||FALSE_ALARM||ACCIDENTAL_ALARM", "NOEMERG||FALSE_ALARM||BOMB_SCARE",
        "NOEMERG||FALSE_ALARM||INTENTIONAL_FALSE_ALARM", "NOEMERG||FALSE_ALARM||MALFUNCTIONING_ALARM",
        "NOEMERG||FALSE_ALARM||OTHER_FALSE_CALL",
    ],
    "NOEMERG – Good Intent": [
        "NOEMERG||GOOD_INTENT||CONTROLLED_BURNING_AUTHORIZED",
        "NOEMERG||GOOD_INTENT||INVESTIGATE_HAZARDOUS_RELEASE",
        "NOEMERG||GOOD_INTENT||NO_INCIDENT_FOUND_LOCATION_ERROR",
        "NOEMERG||GOOD_INTENT||SMOKE_FROM_NONHOSTILE_SOURCE",
    ],
    "PUBSERV – Alarms (Non-Medical)": [
        "PUBSERV||ALARMS_NONMED||CO_ALARM", "PUBSERV||ALARMS_NONMED||FIRE_ALARM",
        "PUBSERV||ALARMS_NONMED||GAS_ALARM", "PUBSERV||ALARMS_NONMED||OTHER_ALARM",
    ],
    "PUBSERV – Citizen Assist": [
        "PUBSERV||CITIZEN_ASSIST||CITIZEN_ASSIST_SERVICE_CALL",
        "PUBSERV||CITIZEN_ASSIST||LIFT_ASSIST",
        "PUBSERV||CITIZEN_ASSIST||LOST_PERSON",
        "PUBSERV||CITIZEN_ASSIST||PERSON_IN_DISTRESS",
    ],
    "PUBSERV – Disaster / Weather": [
        "PUBSERV||DISASTER_WEATHER||DAMAGE_ASSESSMENT",
        "PUBSERV||DISASTER_WEATHER||WEATHER_RESPONSE",
    ],
    "PUBSERV – Other": [
        "PUBSERV||OTHER||DAMAGED_HYDRANT", "PUBSERV||OTHER||MOVE_UP", "PUBSERV||OTHER||STANDBY",
    ],
    "RESCUE – Outside": [
        "RESCUE||OUTSIDE||BACKCOUNTRY_RESCUE", "RESCUE||OUTSIDE||CONFINED_SPACE_RESCUE",
        "RESCUE||OUTSIDE||EXTRICATION_ENTRAPPED", "RESCUE||OUTSIDE||HIGH_ANGLE_RESCUE",
        "RESCUE||OUTSIDE||LIMITED_NO_ACCESS", "RESCUE||OUTSIDE||LOW_ANGLE_RESCUE",
        "RESCUE||OUTSIDE||STEEP_ANGLE_RESCUE", "RESCUE||OUTSIDE||TRENCH",
    ],
    "RESCUE – Structure": [
        "RESCUE||STRUCTURE||BUILDING_STRUCTURE_COLLAPSE", "RESCUE||STRUCTURE||CONFINED_SPACE_RESCUE",
        "RESCUE||STRUCTURE||ELEVATOR_ESCALATOR_RESCUE", "RESCUE||STRUCTURE||EXTRICATION_ENTRAPPED",
    ],
    "RESCUE – Transportation": [
        "RESCUE||TRANSPORTATION||AVIATION_COLLISION_CRASH",
        "RESCUE||TRANSPORTATION||AVIATION_STANDBY",
        "RESCUE||TRANSPORTATION||MOTOR_VEHICLE_EXTRICATION_ENTRAPPED",
        "RESCUE||TRANSPORTATION||TRAIN_RAIL_COLLISION_DERAILMENT",
    ],
    "RESCUE – Water": [
        "RESCUE||WATER||PERSON_IN_WATER_STANDING",
        "RESCUE||WATER||PERSON_IN_WATER_SWIFTWATER",
        "RESCUE||WATER||WATERCRAFT_IN_DISTRESS",
    ],
}


ALL_TYPES = sorted({t for types in INCIDENT_TYPE_MAP.values() for t in types})

THIS_YEAR = datetime.now().year
DATE_RANGE_OPTIONS = [
    "All Records",
    "Occurred Today",
    "Occurred Yesterday",
    "Occurred Last 7 Days",
    "Occurred Last 30 Days",
    "Occurred Last 90 Days",
    f"Occurred in {THIS_YEAR}",
    f"Occurred in {THIS_YEAR - 1}",
    "Custom Date Range",
]


def _ask_date(label):
    while True:
        s = input(f"{label} (YYYY-MM-DD): ").strip()
        try:
            return datetime.strptime(s, "%Y-%m-%d").date()
        except ValueError:
            print("  Invalid format, use YYYY-MM-DD.")


def prompt_incident_types():
    cats = list(INCIDENT_TYPE_MAP.keys())
    print("\n── Incident Type Filter (optional) ──────────────")
    for i, cat in enumerate(cats, 1):
        print(f"  {i:>2}. {cat} ({len(INCIDENT_TYPE_MAP[cat])} types)")
    print("Enter category numbers and/or exact type codes, separated by commas")
    print("  e.g.  1, 3   or   FIRE||STRUCTURE_FIRE||CHIMNEY_FIRE, 9")
    print("  Type ? and a category number to list its types (e.g. ?3), ?? to list every type.")
    print("  Leave blank for ALL incident types.")
    while True:
        raw = input("Incident types [all]: ").strip()
        if not raw:
            return None
        if raw.startswith("?"):
            arg = raw[1:].strip()
            if arg == "?":
                for t in ALL_TYPES:
                    print(f"    {t}")
            elif arg.isdigit() and 1 <= int(arg) <= len(cats):
                for t in INCIDENT_TYPE_MAP[cats[int(arg) - 1]]:
                    print(f"    {t}")
            else:
                print("  Use ?<category number> or ??")
            continue

        chosen, bad = [], []
        for part in (p.strip() for p in raw.split(",")):
            if not part:
                continue
            if part.isdigit() and 1 <= int(part) <= len(cats):
                chosen.extend(INCIDENT_TYPE_MAP[cats[int(part) - 1]])
            elif part.upper() in ALL_TYPES:
                chosen.append(part.upper())
            else:
                bad.append(part)
        if bad:
            print(f"  Not recognized: {', '.join(bad)} — try again (?? lists every type).")
            continue
        chosen = list(dict.fromkeys(chosen))
        print(f"  ✓ {len(chosen)} type(s) selected")
        return chosen


def prompt_config():
    print("\n" + "=" * 60)
    print("  NERIS Incidents Lite — Configuration")
    print("=" * 60)

    print("\n── 1. Credentials ───────────────────────────────")
    username = input("NERIS Email: ").strip()
    print("NERIS Password (note: characters will be visible):")
    password = input("> ").strip()
    if not username or not password:
        sys.exit("✗ Email and password are both required.")

    print("\n── 2. Query Parameters ──────────────────────────")
    while True:
        state_code = input("State Code (two letters, e.g. MI): ").strip().upper()
        if re.fullmatch(r"[A-Z]{2}", state_code):
            break
        print("  Please enter a two-letter state code.")
    entity_id = input("NERIS Entity ID (optional — blank for all departments in the state): ").strip() or None

    incident_types = prompt_incident_types()

    print("\n── 4. Date Range ────────────────────────────────")
    print("(based on when incidents OCCURRED — dispatch call_create — not when they were modified)")
    for i, opt in enumerate(DATE_RANGE_OPTIONS, 1):
        print(f"  {i}. {opt}")
    while True:
        choice = input("Select option [1]: ").strip() or "1"
        if choice.isdigit() and 1 <= int(choice) <= len(DATE_RANGE_OPTIONS):
            range_type = DATE_RANGE_OPTIONS[int(choice) - 1]
            break
        print("  Invalid selection, try again.")

    start_date = end_date = None
    if range_type == "Custom Date Range":
        while True:
            start_date = _ask_date("Start date")
            end_date   = _ask_date("End date")
            if start_date <= end_date:
                break
            print("  Start date must be on or before end date.")

    return dict(username=username, password=password, state_code=state_code,
                entity_id=entity_id, incident_types=incident_types,
                range_type=range_type, start_date=start_date, end_date=end_date)


def calculate_date_range(range_type, start_date=None, end_date=None):
    now   = datetime.now(tz=timezone.utc)
    today = now.replace(hour=0, minute=0, second=0, microsecond=0)

    if range_type == "All Records":
        return None, None
    elif range_type == "Occurred Today":
        return today, None
    elif range_type == "Occurred Yesterday":
        yesterday = today - timedelta(days=1)
        return yesterday, yesterday.replace(hour=23, minute=59, second=59)
    elif range_type == "Occurred Last 7 Days":
        return today - timedelta(days=7), None
    elif range_type == "Occurred Last 30 Days":
        return today - timedelta(days=30), None
    elif range_type == "Occurred Last 90 Days":
        return today - timedelta(days=90), None
    elif range_type.startswith("Occurred in "):
        year = int(range_type.split()[-1])
        return (
            datetime(year, 1,  1,  0,  0,  0, tzinfo=timezone.utc),
            datetime(year, 12, 31, 23, 59, 59, tzinfo=timezone.utc),
        )
    elif range_type == "Custom Date Range" and start_date and end_date:
        return (
            datetime.combine(start_date, datetime.min.time(), tzinfo=timezone.utc),
            datetime.combine(end_date, datetime.max.time().replace(microsecond=0), tzinfo=timezone.utc),
        )
    return None, None


_log = {"path": None, "file": None, "writer": None, "lock": threading.Lock()}


def _init_request_log(state_code):
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    path = DOCS_DIR / f"neris_request_log_{state_code}_{datetime.now():%Y%m%d_%H%M%S}.csv"
    f = open(path, "w", newline="", encoding="utf-8")
    w = csv.writer(f)
    w.writerow(["timestamp_utc", "method", "path", "neris_id_entity",
                "status_code", "x_request_id", "attempt"])
    f.flush()
    _log.update(path=str(path), file=f, writer=w)
    return str(path)


def _log_request(method, url, resp, attempt):
    if _log["writer"] is None:
        return
    try:
        parts = urlsplit(url)
        row = [datetime.now(timezone.utc).isoformat(), method.upper(), parts.path,
               (parse_qs(parts.query).get("neris_id_entity") or [""])[0],
               resp.status_code if resp is not None else "",
               resp.headers.get("x-request-id", "") if resp is not None else "",
               attempt]
        with _log["lock"]:
            _log["writer"].writerow(row)
            _log["file"].flush()
    except Exception as e:
        print(f"  ⚠ (non-fatal) could not write request log row: {e}")


BASE_URL = "https://api.neris.fsri.org/v1"

_auth_lock = threading.Lock()
_patched = {"done": False}
_last_reauth_at = {"t": 0.0}
_blocked = {"msg": None}
_MFA_INPUT_TIMEOUT_SECONDS = 600
_RATE_LIMITS = [(20, 5.0)]
_rate_lock = threading.Lock()
_request_times = [deque() for _ in _RATE_LIMITS]


class NerisPermissionDenied(RuntimeError):
    pass


class NerisAccessBlocked(RuntimeError):
    pass


def _is_permission_denied(res):
    import requests
    if not isinstance(res, requests.Response) or res.status_code != 403:
        return False
    try:
        body = res.json()
    except ValueError:
        return False
    return isinstance(body, dict) and "detail" in body


def _looks_like_auth_drop(res):
    import requests
    if not isinstance(res, requests.Response):
        return False
    if res.status_code == 401:
        return True
    return res.status_code == 403 and not _is_permission_denied(res)


def _input_with_timeout(prompt, timeout_seconds):
    import queue
    q = queue.Queue(maxsize=1)

    def _reader():
        try:
            q.put(input(prompt))
        except Exception:
            pass

    threading.Thread(target=_reader, daemon=True).start()
    try:
        return q.get(timeout=timeout_seconds)
    except queue.Empty:
        raise TimeoutError(f"No MFA code entered within {timeout_seconds}s — "
                           f"requesting a fresh code")


def _blocked_message(res, method, path):
    rid = res.headers.get("x-request-id", "") if res is not None else ""
    return (f"NERIS rejected a freshly issued token (HTTP {res.status_code}) for "
            f"{method.upper()} {path}"
            + (f" [x-request-id {rid}]" if rid else "")
            + ". The login worked, so this isn't an expired token: the account "
              "most likely isn't authorized for this data (role or entity-set "
              "attachment), or the gateway is blocking it. Logging in again won't help.")


def _install_auth_patch():
    if _patched["done"]:
        return
    import requests
    from neris_api_client.client import _NerisApiClient
    from neris_api_client.config import TokenSet, GrantType

    def _safe_update_auth(self):
        if self.tokens.expires_at > datetime.now() + timedelta(seconds=120):
            return

        token_url = f"{self.config.base_url}/token"
        form = {"Content-Type": "application/x-www-form-urlencoded"}

        def _store(got):
            self.tokens = TokenSet(
                access_token=got["access_token"],
                refresh_token=got.get("refresh_token") or "",
                expires_at=datetime.now() + timedelta(seconds=int(got.get("expires_in") or 3600)),
            )
            _last_reauth_at["t"] = time.monotonic()

        if self.tokens.refresh_token:
            res = self._session.post(token_url, headers=form, timeout=30,
                                     data={"grant_type": "refresh_token",
                                           "refresh_token": self.tokens.refresh_token})
            if res.status_code == 200:
                _store(res.json())
                return
            print(f"\n  (token refresh returned HTTP {res.status_code} — logging in again)")

        res = self._session.post(token_url, headers=form, timeout=30,
                                 data={"grant_type": GrantType.PASSWORD,
                                       "username": self.config.username,
                                       "password": self.config.password})
        while True:
            try:
                got = res.json()
            except ValueError:
                got = {}
            if res.status_code == 200:
                _store(got)
                return
            if res.status_code == 202 and got.get("challenge_name"):
                print("\n" + "=" * 60)
                print("  CHECK YOUR EMAIL FOR THE MFA CODE")
                print(f"  (waits up to {_MFA_INPUT_TIMEOUT_SECONDS // 60} minutes, then asks for a fresh code)")
                print("=" * 60)
                code = _input_with_timeout(f"Provide MFA code for {got['challenge_name']}: ",
                                           _MFA_INPUT_TIMEOUT_SECONDS).strip()
                res = self._session.post(token_url, headers=form, timeout=30, data={
                    "grant_type": got["challenge_name"],
                    "username": self.config.username,
                    "session": got.get("session", ""),
                    got["challenge_name"]: code,
                })
                continue
            _last_reauth_at["t"] = time.monotonic()
            self.tokens = TokenSet(access_token="", refresh_token="", expires_at=datetime.min)
            raise RuntimeError(f"Login failed: token endpoint returned HTTP {res.status_code} - "
                               f"{res.text[:300]!r}")

    _NerisApiClient._update_auth = _safe_update_auth

    def _send(self, token, method, path, data=None, params=None, model=None):
        if self.config.validate and model:
            if isinstance(data, str):
                data = model.model_validate_json(data).model_dump(mode="json", by_alias=True)
            if isinstance(data, dict):
                data = model.model_validate(data).model_dump(mode="json", by_alias=True)
        headers = {"Authorization": f"Bearer {token}"}
        res = getattr(self._session, method)(
            f"{self.config.base_url}{path}", json=data, params=params, headers=headers
        )
        try:
            res.raise_for_status()
        except requests.exceptions.HTTPError:
            return res
        try:
            return res.json()
        except requests.exceptions.JSONDecodeError:
            return res.text

    def _locked_call(self, method, path, *args, **kwargs):
        if _blocked["msg"]:
            raise NerisAccessBlocked(_blocked["msg"])
        res = None
        for attempt in (1, 2):
            with _auth_lock:
                try:
                    self._update_auth()
                except Exception as e:
                    _blocked["msg"] = _blocked["msg"] or f"Could not log in again: {e}"
                    raise NerisAccessBlocked(_blocked["msg"])
                token_used = self.tokens.access_token
            res = _send(self, token_used, method, path, *args, **kwargs)
            if not _looks_like_auth_drop(res):
                return res
            if attempt == 2:
                with _auth_lock:
                    _blocked["msg"] = _blocked["msg"] or _blocked_message(res, method, path)
                raise NerisAccessBlocked(_blocked["msg"])
            print(f"  (NERIS rejected the token with HTTP {res.status_code} on "
                  f"{method.upper()} {path} — renewing it and retrying once)")
            wait = False
            with _auth_lock:
                if self.tokens.access_token != token_used:
                    pass
                elif time.monotonic() - _last_reauth_at["t"] < 5.0:
                    wait = True
                else:
                    self.tokens.expires_at = datetime.min
                    try:
                        self._update_auth()
                    except Exception as e:
                        _blocked["msg"] = _blocked["msg"] or f"Could not log in again: {e}"
                        raise NerisAccessBlocked(_blocked["msg"])
            if wait:
                time.sleep(1.5)
        return res

    _NerisApiClient._call = _locked_call

    orig_request = _NerisApiClient._session.request

    def _wait_for_rate_budget():
        while True:
            with _rate_lock:
                now = time.monotonic()
                sleep_for = 0.0
                for (max_req, window), times in zip(_RATE_LIMITS, _request_times):
                    while times and now - times[0] > window:
                        times.popleft()
                    if len(times) >= max_req:
                        sleep_for = max(sleep_for, window - (now - times[0]) + 0.05)
                if sleep_for <= 0:
                    for times in _request_times:
                        times.append(now)
                    return
            time.sleep(sleep_for)

    def _throttled_request(method, url, *args, **kwargs):
        kwargs.setdefault("timeout", 60)
        for attempt in range(1, 7):
            _wait_for_rate_budget()
            resp = orig_request(method, url, *args, **kwargs)
            _log_request(method, url, resp, attempt)
            if resp.status_code != 429 or attempt == 6:
                return resp
            wait = min(120, 5 * 2 ** (attempt - 1))
            print(f"  ⚠ Rate limited (HTTP 429) — waiting {wait}s (attempt {attempt}/6)")
            time.sleep(wait)
        return resp

    _NerisApiClient._session.request = _throttled_request
    _patched["done"] = True


def _unwrap(res, context):
    import requests
    if isinstance(res, requests.Response):
        if _is_permission_denied(res):
            raise NerisPermissionDenied(f"{context}: HTTP 403 {res.text[:200]}")
        raise RuntimeError(f"{context} failed: HTTP {res.status_code} - {res.text[:300]!r}")
    return res


def authenticate(username, password):
    from neris_api_client import NerisApiClient, Config
    from neris_api_client.config import GrantType
    _install_auth_patch()
    print("\nConnecting to NERIS API...")
    client = NerisApiClient(Config(
        base_url=BASE_URL,
        grant_type=GrantType.PASSWORD,
        username=username,
        password=password,
    ))
    for attempt in range(1, 4):
        try:
            _unwrap(client.health(), "auth check (health)")
            print("✓ Authentication successful!")
            return client
        except TimeoutError as e:
            if attempt < 3:
                print(f"\n⚠ {e}")
                continue
            sys.exit(f"✗ Authentication failed: {e}")
        except Exception as e:
            if "HTTP 50" in str(e) and attempt < 3:
                print(f"⚠ Transient server error ({e}) — retrying in {attempt * 5}s")
                time.sleep(attempt * 5)
                continue
            sys.exit(f"✗ Authentication failed: {e}")


_fetch_status = {"incomplete": False, "error": None}


def _with_retries(fn, label, attempts=3):
    for attempt in range(1, attempts + 1):
        try:
            return fn()
        except (NerisPermissionDenied, NerisAccessBlocked):
            raise
        except Exception as e:
            if attempt == attempts:
                raise
            print(f"\n  ⚠ {label} failed ({e}) — retrying in {5 * attempt}s")
            time.sleep(5 * attempt)


def get_state_incidents(client, state_code, call_create_start=None, call_create_end=None,
                        neris_id_entity=None, incident_types=None, page_size=100):
    all_incidents, next_cursor, page = [], None, 0
    _fetch_status.update(incomplete=False, error=None)

    print(f"\nRetrieving incidents for state: {state_code}")
    if neris_id_entity:
        print(f"Filter: neris_id_entity = {neris_id_entity}")
    if call_create_start:
        print(f"Filter: call_create >= {call_create_start:%Y-%m-%d %H:%M:%S}")
    if call_create_end:
        print(f"Filter: call_create <= {call_create_end:%Y-%m-%d %H:%M:%S}")
    print(f"Filter: incident_types = {len(incident_types)} type(s)" if incident_types
          else "Filter: All incident types (no type filter)")
    if not call_create_start and not call_create_end:
        print("Filter: All records (no date filter)")

    while True:
        page += 1
        print(f"  Page {page}... ", end="", flush=True)
        kwargs = {"state": state_code, "page_size": page_size, "geo_format": "geojson"}
        if call_create_start: kwargs["call_create_start"] = call_create_start
        if call_create_end:   kwargs["call_create_end"]   = call_create_end
        if neris_id_entity:   kwargs["neris_id_entity"]   = neris_id_entity
        if incident_types:    kwargs["incident_types"]    = list(incident_types)
        if next_cursor:       kwargs["cursor"]            = next_cursor

        try:
            res = _with_retries(
                lambda: _unwrap(client.list_incidents(**kwargs), f"list_incidents(page {page})"),
                f"page {page}")
        except NerisAccessBlocked:
            _fetch_status.update(incomplete=True, error=_blocked["msg"])
            raise
        except Exception as e:
            _fetch_status.update(incomplete=True, error=str(e))
            print(f"\n✗ Page {page} failed: {e}")
            print(f"⚠ STOPPING EARLY — only {len(all_incidents)} incident(s) from "
                  f"{page - 1} page(s) were retrieved. Results are INCOMPLETE.")
            break

        incidents = res.get("incidents", []) or []
        if not incidents:
            print("empty page")
            break
        all_incidents.extend(incidents)
        print(f"retrieved {len(incidents)} incidents (Total: {len(all_incidents)})")
        next_cursor = res.get("next_cursor")
        if not next_cursor:
            print("✓ No more pages")
            break

    print(f"\n{'=' * 50}")
    print(f"Total incidents retrieved: {len(all_incidents)}")
    if _fetch_status["incomplete"]:
        print("⚠ INCOMPLETE — the fetch stopped early (see errors above)")
    print("=" * 50)
    return all_incidents


def lookup_department_names(client, dept_ids, max_workers=8):
    print(f"\nLooking up {len(dept_ids)} department names...")
    names = {}

    def _one(did):
        try:
            entity = _with_retries(lambda: _unwrap(client.get_entity(did), f"get_entity({did})"),
                                   f"get_entity({did})")
            return did, entity.get("name", "") if isinstance(entity, dict) else ""
        except NerisAccessBlocked:
            return did, ""
        except Exception as e:
            print(f"  ⚠ Could not look up department name for {did}: {e}")
            return did, ""

    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        for i, (did, name) in enumerate(ex.map(_one, sorted(dept_ids)), 1):
            names[did] = name
            if i % 25 == 0:
                print(f"  {i}/{len(dept_ids)} looked up...")
    print("✓ Department lookup complete")
    return names


def build_address(locs):
    if isinstance(locs, dict):
        locs = [locs]
    locs = [l for l in (locs or []) if isinstance(l, dict) and l]
    if not locs:
        return ""

    def first(key):
        for l in locs:
            v = l.get(key)
            if v not in (None, ""):
                return str(v).strip()
        return ""

    def house_number(l):
        num = l.get("complete_number")
        if num not in (None, ""):
            return str(num).strip()
        num = f"{l.get('number_prefix') or ''}{l.get('number') or ''}"
        if l.get("number_suffix"):
            num = f"{num} {l['number_suffix']}"
        return num.strip()

    number = next((n for n in (house_number(l) for l in locs) if n), "")
    street = " ".join(p for p in [number, first("street_prefix_direction"), first("street"),
                                  first("street_postfix"), first("street_postfix_direction")] if p)
    unit = first("unit_value")
    if unit:
        street = f"{street} {first('unit_prefix') or 'Apt'} {unit}".strip()
    city_state = ", ".join(p for p in [first("incorporated_municipality"), first("state")] if p)
    zip_ = first("postal_code")
    if zip_:
        city_state = f"{city_state} {zip_}".strip()
    return ", ".join(p for p in [street, city_state] if p)


def get_xy_from_point(point_obj):
    if not point_obj or not isinstance(point_obj, dict):
        return "", ""
    coords = (point_obj.get("geometry") or point_obj).get("coordinates", [])
    if coords and len(coords) >= 2:
        return coords[0], coords[1]
    return "", ""


def count_casualties(casualty_rescues):
    counts = {"ff_fatal": 0, "ff_injured": 0, "ff_rescue": 0,
              "civilian_fatal": 0, "civilian_injured": 0, "civilian_rescue": 0}
    if not casualty_rescues or not isinstance(casualty_rescues, list):
        return counts
    outcome = {"INJURED_FATAL": "fatal", "INJURED_NONFATAL": "injured", "UNINJURED": "rescue"}
    for rec in casualty_rescues:
        if not isinstance(rec, dict):
            continue
        role = {"FF": "ff", "CIVILIAN": "civilian"}.get((rec.get("type") or "").upper())
        inj  = (((rec.get("casualty") or {}).get("injury_or_noninjury") or {}).get("type") or "").upper()
        if role and inj in outcome:
            counts[f"{role}_{outcome[inj]}"] += 1
    return counts


def export_lite(incidents, client, state_code):
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.cell import WriteOnlyCell
    try:
        from openpyxl.cell.cell import KNOWN_TYPES as _XL_TYPES
    except ImportError:
        _XL_TYPES = (str, int, float, bool, datetime, type(None))

    if not incidents:
        print("⚠ No incidents to export.")
        return None

    dept_ids = {(inc.get("base") or {}).get("department_neris_id") for inc in incidents} - {None, ""}
    dept_names = lookup_department_names(client, dept_ids)

    def join_list(lst):
        if not lst or not isinstance(lst, list): return ""
        return "||".join(str(x) for x in lst if x)

    def extract_list_field(lst, key):
        if not lst or not isinstance(lst, list): return ""
        return "||".join(str(item.get(key, "")) for item in lst
                         if isinstance(item, dict) and item.get(key))

    rows = []
    for inc in incidents:
        base       = inc.get("base", {}) or {}
        dispatch   = inc.get("dispatch", {}) or {}
        status_obj = inc.get("incident_status", {}) or {}
        dept_id    = base.get("department_neris_id", "")

        inc_types = inc.get("incident_types", []) or []
        type_list = [it.get("type", "") for it in inc_types if isinstance(it, dict) and it.get("type")]

        base_loc = base.get("location", {}) or {}
        base_geo = base_loc.get("geocoded_location", {}) or {}
        disp_loc = dispatch.get("location", {}) or {}
        disp_geo = disp_loc.get("geocoded_location", {}) or {}

        base_point = base.get("point") or {}
        if base_point:
            point_src = base_point
        elif base_geo:
            point_src = base_geo.get("point") or {}
        elif disp_geo:
            point_src = disp_geo.get("point") or {}
        else:
            point_src = {}

        addr_chain = [base_geo, base_loc, disp_geo, disp_loc]
        address = build_address(addr_chain)
        county  = next((l.get("county") for l in addr_chain if l.get("county")), "")
        x, y    = get_xy_from_point(point_src)

        aids  = inc.get("aids", []) or []
        units = inc.get("unit_responses", []) or []
        an    = (inc.get("actions_tactics", {}) or {}).get("action_noaction", {}) or {}
        cas   = count_casualties(inc.get("casualty_rescues", []))
        loc_use  = base.get("location_use", {}) or {}
        fire     = inc.get("fire_detail", {}) or {}
        fire_loc = fire.get("location_detail", {}) or {}

        smoke_pres = (inc.get("smoke_alarm", {}) or {}).get("presence", {}) or {}
        fire_pres  = (inc.get("fire_alarm",  {}) or {}).get("presence", {}) or {}
        other_pres = (inc.get("other_alarm", {}) or {}).get("presence", {}) or {}
        fs_pres    = (inc.get("fire_suppression", {}) or {}).get("presence", {}) or {}
        fs_supp    = fs_pres.get("suppression_types", []) or []
        fs_supp_str = (extract_list_field(fs_supp, "type") if fs_supp and isinstance(fs_supp[0], dict)
                       else join_list(fs_supp))
        med_o2_pres = (inc.get("medical_oxygen_hazard",    {}) or {}).get("presence", {}) or {}
        cfs_pres    = (inc.get("cooking_fire_suppression", {}) or {}).get("presence", {}) or {}

        rows.append({
            "Incident NERIS ID":                      inc.get("neris_id", ""),
            "Department/Entity ID":                   dept_id,
            "Department Name":                        dept_names.get(dept_id, ""),
            "Incident Number":                        base.get("incident_number", ""),
            "Call Date":                              dispatch.get("call_create", ""),
            "Incident Status":                        status_obj.get("status", ""),
            "Address":                                address,
            "County":                                 county,
            "First Incident Type":                    type_list[0] if len(type_list) > 0 else "",
            "Second Incident Type":                   type_list[1] if len(type_list) > 1 else "",
            "Third Incident Type":                    type_list[2] if len(type_list) > 2 else "",
            "Aid Types":                              extract_list_field(aids, "aid_type"),
            "Aid Directions":                         extract_list_field(aids, "aid_direction"),
            "Aid Department IDs":                     extract_list_field(aids, "department_neris_id"),
            "Non-FD Aid Types":                       extract_list_field(inc.get("nonfd_aids", []), "type"),
            "Incident Modifiers":                     extract_list_field(inc.get("special_modifiers", []), "type"),
            "Total Units":                            len(units),
            "Total Staffing":                         sum(u.get("staffing", 0) or 0 for u in units if isinstance(u, dict)),
            "Exposure Count":                         len(inc.get("exposures", []) or []),
            "Actions/Tactics Type":                   an.get("type", ""),
            "Actions/Tactics Actions":                join_list(an.get("actions", [])),
            "FF Fatal Count":                         cas["ff_fatal"],
            "FF Injured Count":                       cas["ff_injured"],
            "FF Rescue (Uninjured) Count":            cas["ff_rescue"],
            "Civilian Fatal Count":                   cas["civilian_fatal"],
            "Civilian Injured Count":                 cas["civilian_injured"],
            "Civilian Rescue (Uninjured) Count":      cas["civilian_rescue"],
            "Location Use Type":                      loc_use.get("use_type", ""),
            "Location Secondary Use":                 loc_use.get("secondary_use", ""),
            "Acres Burned":                           fire_loc.get("acres_burned", ""),
            "Smoke Alarm Presence Type":              smoke_pres.get("type", ""),
            "Smoke Alarm Types":                      join_list(smoke_pres.get("alarm_types", [])),
            "Fire Alarm Presence Type":               fire_pres.get("type", ""),
            "Fire Alarm Types":                       join_list(fire_pres.get("alarm_types", [])),
            "Other Alarm Presence Type":              other_pres.get("type", ""),
            "Other Alarm Types":                      join_list(other_pres.get("alarm_types", [])),
            "Fire Suppression Presence Type":         fs_pres.get("type", ""),
            "Fire Suppression Types":                 fs_supp_str,
            "Medical Oxygen Hazard Presence Type":    med_o2_pres.get("type", ""),
            "Cooking Fire Suppression Presence Type": cfs_pres.get("type", ""),
            "Fire Cause":                             fire_loc.get("cause", ""),
            "X (Longitude)":                          x,
            "Y (Latitude)":                           y,
        })

    df = pd.DataFrame(rows).fillna("")
    keep_always = {
        "X (Longitude)", "Y (Latitude)",
        "FF Fatal Count", "FF Injured Count", "FF Rescue (Uninjured) Count",
        "Civilian Fatal Count", "Civilian Injured Count", "Civilian Rescue (Uninjured) Count",
        "Total Units", "Total Staffing", "Exposure Count",
    }
    df = df.drop(columns=[c for c in df.columns if c not in keep_always
                          and df[c].apply(lambda v: v == "" or v == 0 or str(v) == "0").all()])

    print(f"\n  Incidents Lite    : {len(df)} rows, {len(df.columns)} columns")
    illegal = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
    EXCEL_CELL_MAX = 32767

    def sanitize(val):
        if isinstance(val, (list, tuple, set)):
            if not val:
                return None
            val = ("||".join(str(x) for x in val)
                   if all(not isinstance(x, (dict, list, tuple, set)) for x in val)
                   else json.dumps(list(val), default=str))
        elif isinstance(val, dict):
            val = json.dumps(val, default=str) if val else None
        elif isinstance(val, datetime) and val.tzinfo is not None:
            val = val.astimezone(timezone.utc).replace(tzinfo=None)
        try:
            if pd.isna(val):
                return None
        except (TypeError, ValueError):
            pass
        if isinstance(val, str):
            val = illegal.sub("", val)
            if len(val) > EXCEL_CELL_MAX:
                suffix = " ...[truncated]"
                val = val[:EXCEL_CELL_MAX - len(suffix)] + suffix
        elif val is not None and not isinstance(val, _XL_TYPES):
            val = str(val)
        return val

    fill  = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font  = Font(color="FFFFFF", bold=True, size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def write_sheet(wb, frame, name):
        ws = wb.create_sheet(name)
        headers = list(frame.columns)
        widths = [len(str(h)) for h in headers]
        for row in frame.head(50).itertuples(index=False, name=None):
            for i, v in enumerate(row):
                v = sanitize(v)
                if v:
                    widths[i] = max(widths[i], len(str(v)))
        for col_idx, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(w + 2, 12), 60)
        ws.freeze_panes = "A2"
        header_cells = []
        for h in headers:
            c = WriteOnlyCell(ws, value=h)
            c.fill, c.font, c.alignment = fill, font, align
            header_cells.append(c)
        ws.append(header_cells)
        for i, row in enumerate(frame.itertuples(index=False, name=None), 1):
            ws.append([sanitize(v) for v in row])
            if i % 20000 == 0:
                print(f"    {name}: {i:,}/{len(frame):,} rows written...")

    wb = Workbook(write_only=True)
    write_sheet(wb, df, "Incidents Lite")

    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    filepath = str(DOCS_DIR / f"NERIS Incidents Lite {state_code} {datetime.now():%Y-%m-%d}.xlsx")
    wb.save(filepath)

    print(f"\n✓ Exported: {filepath}")
    print(f"  Incidents          : {len(incidents)}")
    print("  Sheet              : Incidents Lite")
    return filepath


def _print_blocked(msg):
    print("\n" + "!" * 70)
    print("  ✗ STOPPED: NERIS is refusing this account's requests.")
    print(f"    {msg}")
    print("    Check the account's role / entity-set attachment for this state")
    print("    or department. The request log below has the x-request-id(s) to")
    print("    send to the NERIS team.")
    print("!" * 70)


def main():
    ensure_dependencies()
    cfg = prompt_config()

    log_path = _init_request_log(cfg["state_code"])
    print(f"\n✓ Logging every request's x-request-id to: {log_path}")

    client = authenticate(cfg["username"], cfg["password"])

    print(f"\nState         : {cfg['state_code']}")
    if cfg["entity_id"]:
        print(f"Entity ID     : {cfg['entity_id']}")
    print(f"Date Range    : {cfg['range_type']}")
    print(f"Incident Types: {len(cfg['incident_types'])} selected" if cfg["incident_types"]
          else "Incident Types: All (no filter)")

    start_filter, end_filter = calculate_date_range(cfg["range_type"], cfg["start_date"], cfg["end_date"])

    try:
        incidents = get_state_incidents(
            client, cfg["state_code"],
            call_create_start=start_filter, call_create_end=end_filter,
            neris_id_entity=cfg["entity_id"], incident_types=cfg["incident_types"],
        )
        if not incidents:
            print("\n⚠ No incidents found.")
        else:
            print("\nBuilding Incidents Lite export...")
            export_lite(incidents, client, cfg["state_code"])
    except NerisAccessBlocked as e:
        _print_blocked(e)
    except KeyboardInterrupt:
        print("\n✗ Cancelled.")
    except Exception as e:
        print(f"\n✗ Error: {e}")
        traceback.print_exc()

    if _blocked["msg"] and _fetch_status["error"] != _blocked["msg"]:
        _print_blocked(_blocked["msg"])
    if _fetch_status["incomplete"] and _fetch_status["error"] != _blocked["msg"]:
        print("\n" + "!" * 70)
        print("  ⚠ WARNING: the incident fetch stopped early — the export (if any)")
        print("    contains PARTIAL results, not the full set. Last error:")
        print(f"    {_fetch_status['error']}")
        print("    Re-run the script to get the complete data.")
        print("!" * 70)

    if _log["file"]:
        _log["file"].flush()
        print(f"\n  Request trace log : {_log['path']}")

    print("\n" + "=" * 70)
    print("✗ PROCESS STOPPED (see above)" if _blocked["msg"] and _fetch_status["incomplete"]
          else "✓ PROCESS COMPLETE")
    print("=" * 70)


if __name__ == "__main__":
    main()
